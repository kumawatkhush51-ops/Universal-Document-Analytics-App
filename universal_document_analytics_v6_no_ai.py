import io
import json
import math
import re
import textwrap
from datetime import datetime

import numpy as np
import pandas as pd
import streamlit as st

# Optional document readers. The app shows a friendly message if a reader is unavailable.
try:
    from docx import Document
except Exception:
    Document = None

try:
    from pypdf import PdfReader
except Exception:
    PdfReader = None

try:
    from pptx import Presentation
except Exception:
    Presentation = None


# ------------------------------------------------------------
# PAGE CONFIG
# ------------------------------------------------------------
st.set_page_config(
    page_title="Universal Document Analytics",
    page_icon="📊",
    layout="wide",
    initial_sidebar_state="expanded",
)

st.markdown(
    """
    <style>
    .main-title {
        font-size: 42px;
        font-weight: 800;
        margin-bottom: 4px;
    }
    .subtitle {
        color: #9aa4b2;
        font-size: 17px;
        margin-bottom: 24px;
    }
    .metric-card {
        border: 1px solid #30343d;
        border-radius: 14px;
        padding: 18px;
        min-height: 130px;
        background: rgba(255,255,255,0.015);
    }
    .metric-label {
        font-size: 14px;
        color: #b8c0cc;
        margin-bottom: 8px;
    }
    .metric-value {
        font-size: 28px;
        font-weight: 750;
        line-height: 1.25;
        overflow-wrap: anywhere;
        word-break: break-word;
    }
    .insight-card {
        border: 1px solid #30343d;
        border-radius: 14px;
        padding: 18px;
        min-height: 150px;
        background: rgba(255,255,255,0.015);
    }
    .insight-title {
        font-size: 19px;
        font-weight: 700;
        margin-bottom: 10px;
    }
    .small-muted {
        color: #9aa4b2;
        font-size: 13px;
    }
    .success-box {
        border: 1px solid #235f3a;
        background: rgba(35,95,58,0.18);
        padding: 14px 16px;
        border-radius: 12px;
    }
    .warning-box {
        border: 1px solid #705c1c;
        background: rgba(112,92,28,0.16);
        padding: 14px 16px;
        border-radius: 12px;
    }
    </style>
    """,
    unsafe_allow_html=True,
)


# ------------------------------------------------------------
# GENERAL HELPERS
# ------------------------------------------------------------
def clean_column_name(value):
    """Make a column safe and readable without destroying its meaning."""
    text = str(value).strip()
    text = re.sub(r"\s+", " ", text)
    return text if text else "Unnamed"


def clean_dataframe(df):
    if df is None:
        return pd.DataFrame()

    df = df.copy()

    # Make duplicate column names unique.
    seen = {}
    new_columns = []
    for col in df.columns:
        base = clean_column_name(col)
        count = seen.get(base, 0)
        seen[base] = count + 1
        new_columns.append(base if count == 0 else f"{base}_{count + 1}")
    df.columns = new_columns

    # Remove completely empty rows/columns.
    df = df.dropna(axis=0, how="all")
    df = df.dropna(axis=1, how="all")
    return df.reset_index(drop=True)


def read_uploaded_file(uploaded_file):
    """
    Returns:
        {
          "kind": "table" | "text",
          "name": str,
          "dataframe": pd.DataFrame | None,
          "text": str,
          "error": str | None
        }
    """
    name = uploaded_file.name
    suffix = Path(name).suffix.lower()
    raw = uploaded_file.getvalue()

    result = {
        "kind": "text",
        "name": name,
        "dataframe": None,
        "text": "",
        "error": None,
    }

    try:
        if suffix in {".xlsx", ".xlsm"}:
            df = pd.read_excel(io.BytesIO(raw))
            result["kind"] = "table"
            result["dataframe"] = clean_dataframe(df)
            return result

        if suffix == ".xls":
            df = pd.read_excel(io.BytesIO(raw), engine="xlrd")
            result["kind"] = "table"
            result["dataframe"] = clean_dataframe(df)
            return result

        if suffix == ".csv":
            # utf-8 first, then common Windows encoding.
            try:
                df = pd.read_csv(io.BytesIO(raw))
            except UnicodeDecodeError:
                df = pd.read_csv(io.BytesIO(raw), encoding="latin1")
            result["kind"] = "table"
            result["dataframe"] = clean_dataframe(df)
            return result

        if suffix == ".json":
            obj = json.loads(raw.decode("utf-8-sig"))
            if isinstance(obj, list):
                df = pd.json_normalize(obj)
            elif isinstance(obj, dict):
                # Prefer a list nested in a dictionary when present.
                list_value = next(
                    (v for v in obj.values() if isinstance(v, list)),
                    None,
                )
                if list_value is not None:
                    df = pd.json_normalize(list_value)
                else:
                    df = pd.json_normalize(obj)
            else:
                raise ValueError("The JSON structure could not be converted to a table.")
            result["kind"] = "table"
            result["dataframe"] = clean_dataframe(df)
            return result

        if suffix in {".txt", ".md", ".log", ".html", ".htm", ".xml", ".rtf"}:
            text = raw.decode("utf-8-sig", errors="replace")
            # Remove basic HTML tags for a cleaner preview.
            if suffix in {".html", ".htm"}:
                text = re.sub(r"<[^>]+>", " ", text)
            text = re.sub(r"\s+", " ", text).strip()
            result["kind"] = "text"
            result["text"] = text
            return result

        if suffix == ".docx":
            if Document is None:
                raise RuntimeError("python-docx is not installed.")
            doc = Document(io.BytesIO(raw))
            paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip()]

            # Include tables in a readable form.
            table_lines = []
            for table in doc.tables:
                for row in table.rows:
                    cells = [cell.text.strip().replace("\n", " ") for cell in row.cells]
                    if any(cells):
                        table_lines.append(" | ".join(cells))

            text = "\n".join(paragraphs + table_lines)
            result["kind"] = "text"
            result["text"] = text.strip()
            return result

        if suffix == ".pdf":
            if PdfReader is None:
                raise RuntimeError("pypdf is not installed.")
            reader = PdfReader(io.BytesIO(raw))
            pages = []
            for page in reader.pages:
                try:
                    pages.append(page.extract_text() or "")
                except Exception:
                    pages.append("")
            result["kind"] = "text"
            result["text"] = "\n".join(pages).strip()
            return result

        if suffix == ".pptx":
            if Presentation is None:
                raise RuntimeError("python-pptx is not installed.")
            prs = Presentation(io.BytesIO(raw))
            chunks = []
            for slide_number, slide in enumerate(prs.slides, start=1):
                chunks.append(f"Slide {slide_number}")
                for shape in slide.shapes:
                    if hasattr(shape, "text") and shape.text.strip():
                        chunks.append(shape.text.strip())
            result["kind"] = "text"
            result["text"] = "\n".join(chunks).strip()
            return result

        raise ValueError(
            f"Unsupported file type: {suffix or 'unknown'}. "
            "Supported: XLSX, XLS, XLSM, CSV, JSON, DOCX, PDF, PPTX, TXT, MD, HTML, XML and RTF."
        )

    except Exception as exc:
        result["error"] = str(exc)
        return result


def normalize_name(name):
    return re.sub(r"[^a-z0-9]+", "", str(name).lower())


def find_column(df, keywords, exclude=None):
    """Find a column using exact/contains matching. Returns the real column name or None."""
    if df is None or df.empty:
        return None

    exclude = set(exclude or [])
    normalized = {col: normalize_name(col) for col in df.columns}

    # Exact normalized match first.
    for keyword in keywords:
        nk = normalize_name(keyword)
        for col, nc in normalized.items():
            if col not in exclude and nc == nk:
                return col

    # Contains match second.
    for keyword in keywords:
        nk = normalize_name(keyword)
        for col, nc in normalized.items():
            if col not in exclude and nk and nk in nc:
                return col

    return None


def numeric_series(df, column):
    if not column or column not in df.columns:
        return pd.Series(dtype="float64")
    return pd.to_numeric(df[column], errors="coerce")


def money(value):
    if value is None or (isinstance(value, float) and math.isnan(value)):
        return "—"
    try:
        value = float(value)
        sign = "-" if value < 0 else ""
        value = abs(value)
        return f"{sign}₹{value:,.2f}"
    except Exception:
        return str(value)


def number(value):
    if value is None or (isinstance(value, float) and math.isnan(value)):
        return "—"
    try:
        value = float(value)
        if value.is_integer():
            return f"{int(value):,}"
        return f"{value:,.2f}"
    except Exception:
        return str(value)


def percent(value):
    if value is None or (isinstance(value, float) and math.isnan(value)):
        return "—"
    try:
        return f"{float(value):,.1f}%"
    except Exception:
        return str(value)


def safe_text(value, max_chars=700):
    text = str(value).replace("\x00", " ").strip()
    return text if len(text) <= max_chars else text[:max_chars].rstrip() + "…"


# ------------------------------------------------------------
# DATA ANALYTICS
# ------------------------------------------------------------
def analyze_table(df):
    df = clean_dataframe(df)
    result = {
        "rows": len(df),
        "columns": len(df.columns),
        "missing": int(df.isna().sum().sum()),
        "quality": 100.0,
        "numeric_cols": [],
        "date_col": None,
        "sales_col": None,
        "quantity_col": None,
        "price_col": None,
        "cost_col": None,
        "profit_col": None,
        "category_col": None,
        "product_col": None,
        "period_col": None,
        "metrics": {},
        "highlights": [],
        "diagnostic": [],
        "predictive": [],
        "prescriptive": [],
    }

    if df.empty:
        return result

    total_cells = max(1, df.shape[0] * df.shape[1])
    result["quality"] = max(0.0, 100.0 * (1 - result["missing"] / total_cells))

    numeric_cols = []
    for col in df.columns:
        converted = pd.to_numeric(df[col], errors="coerce")
        if converted.notna().sum() >= max(2, int(len(df) * 0.5)):
            numeric_cols.append(col)
    result["numeric_cols"] = numeric_cols

    result["sales_col"] = find_column(
        df, ["sales", "revenue", "salesamount", "revenueamount", "turnover"]
    )
    result["quantity_col"] = find_column(
        df, ["quantity", "qty", "quantitysold", "units", "unitsold", "volume"]
    )
    result["price_col"] = find_column(
        df, ["unitprice", "price", "sellingprice", "averageprice"]
    )
    result["cost_col"] = find_column(
        df, ["unitcost", "cost", "costofgoods", "cogs", "costofgoodssold"]
    )
    result["profit_col"] = find_column(
        df, ["profit", "grossprofit", "netprofit", "margin"]
    )
    result["category_col"] = find_column(
        df, ["category", "productcategory", "segment", "department", "type"]
    )
    result["product_col"] = find_column(
        df, ["product", "productname", "item", "itemname", "sku"]
    )
    result["period_col"] = find_column(
        df, ["month", "period", "quarter", "year", "date", "billingdate", "salesdate"]
    )

    # Date detection.
    for col in df.columns:
        if result["date_col"]:
            break
        if pd.api.types.is_datetime64_any_dtype(df[col]):
            result["date_col"] = col
            break
        name = normalize_name(col)
        if any(k in name for k in ["date", "billingdate", "salesdate", "transactiondate"]):
            parsed = pd.to_datetime(df[col], errors="coerce")
            if parsed.notna().sum() >= max(2, int(len(df) * 0.5)):
                result["date_col"] = col

    # Metrics.
    if result["sales_col"]:
        sales = numeric_series(df, result["sales_col"])
        result["metrics"]["total_sales"] = float(sales.sum())
        result["metrics"]["average_sale"] = float(sales.mean()) if sales.notna().any() else None
        result["metrics"]["max_sale"] = float(sales.max()) if sales.notna().any() else None
        result["metrics"]["min_sale"] = float(sales.min()) if sales.notna().any() else None
    else:
        result["metrics"]["total_sales"] = None

    if result["quantity_col"]:
        qty = numeric_series(df, result["quantity_col"])
        result["metrics"]["total_quantity"] = float(qty.sum())
    else:
        result["metrics"]["total_quantity"] = None

    if result["price_col"]:
        prices = numeric_series(df, result["price_col"])
        result["metrics"]["average_price"] = float(prices.mean()) if prices.notna().any() else None
    else:
        result["metrics"]["average_price"] = None

    # Calculated profit when possible.
    if result["profit_col"]:
        profit = numeric_series(df, result["profit_col"])
        result["metrics"]["profit"] = float(profit.sum())
    elif result["sales_col"] and result["cost_col"]:
        sales = numeric_series(df, result["sales_col"])
        costs = numeric_series(df, result["cost_col"])
        result["metrics"]["profit"] = float((sales - costs).sum())
    elif result["sales_col"] and result["quantity_col"] and result["cost_col"] and result["price_col"]:
        sales = numeric_series(df, result["sales_col"])
        qty = numeric_series(df, result["quantity_col"])
        cost = numeric_series(df, result["cost_col"])
        price = numeric_series(df, result["price_col"])
        result["metrics"]["profit"] = float((sales - qty * cost).sum())
    else:
        result["metrics"]["profit"] = None

    if result["metrics"]["total_sales"] is not None and result["metrics"]["profit"] is not None:
        sales_total = result["metrics"]["total_sales"]
        result["metrics"]["margin"] = (
            100 * result["metrics"]["profit"] / sales_total
            if sales_total else None
        )
    else:
        result["metrics"]["margin"] = None

    # Highlights.
    if result["sales_col"]:
        sales = numeric_series(df, result["sales_col"])
        if result["product_col"]:
            grouped = pd.DataFrame({
                "label": df[result["product_col"]].astype(str),
                "value": sales,
            }).dropna()
            if not grouped.empty:
                top = grouped.groupby("label")["value"].sum().sort_values(ascending=False)
                if not top.empty:
                    result["highlights"].append(
                        ("🏆", "Top Product", str(top.index[0]), money(top.iloc[0]))
                    )

        if result["category_col"]:
            grouped = pd.DataFrame({
                "label": df[result["category_col"]].astype(str),
                "value": sales,
            }).dropna()
            if not grouped.empty:
                top = grouped.groupby("label")["value"].sum().sort_values(ascending=False)
                if not top.empty:
                    total = top.sum()
                    share = 100 * top.iloc[0] / total if total else None
                    result["highlights"].append(
                        ("⭐", "Leading Category", str(top.index[0]), percent(share) + " of sales")
                    )

        if result["period_col"]:
            grouped = pd.DataFrame({
                "label": df[result["period_col"]].astype(str),
                "value": sales,
            }).dropna()
            if not grouped.empty:
                top = grouped.groupby("label")["value"].sum().sort_values(ascending=False)
                if not top.empty:
                    result["highlights"].append(
                        ("📅", "Best Period", str(top.index[0]), money(top.iloc[0]))
                    )

    # Diagnostic analysis.
    if result["product_col"] and result["sales_col"]:
        sales = numeric_series(df, result["sales_col"])
        temp = pd.DataFrame({
            "product": df[result["product_col"]].astype(str),
            "sales": sales,
        }).dropna()
        if not temp.empty:
            grouped = temp.groupby("product")["sales"].sum().sort_values(ascending=False)
            if len(grouped) >= 2:
                result["diagnostic"].append(
                    f"{grouped.index[0]} leads sales at {money(grouped.iloc[0])}; "
                    f"{grouped.index[-1]} is lowest at {money(grouped.iloc[-1])}."
                )

    if result["category_col"] and result["sales_col"]:
        sales = numeric_series(df, result["sales_col"])
        temp = pd.DataFrame({
            "category": df[result["category_col"]].astype(str),
            "sales": sales,
        }).dropna()
        if not temp.empty:
            grouped = temp.groupby("category")["sales"].sum().sort_values(ascending=False)
            if len(grouped) >= 2:
                gap = grouped.iloc[0] - grouped.iloc[-1]
                result["diagnostic"].append(
                    f"The gap between the strongest and weakest category is {money(gap)}."
                )

    # Simple predictive trend.
    if result["date_col"] and result["sales_col"]:
        dates = pd.to_datetime(df[result["date_col"]], errors="coerce")
        sales = numeric_series(df, result["sales_col"])
        trend_df = pd.DataFrame({"date": dates, "sales": sales}).dropna().sort_values("date")
        if len(trend_df) >= 3:
            first = trend_df["sales"].iloc[0]
            last = trend_df["sales"].iloc[-1]
            if first != 0:
                change = 100 * (last - first) / abs(first)
                direction = "upward" if change > 0 else "downward" if change < 0 else "flat"
                result["predictive"].append(
                    f"The observed sales trend is {direction}, with a change of {change:.1f}% "
                    f"from the first to the latest available period."
                )
        else:
            result["predictive"].append(
                "There are not enough valid dated sales observations for a reliable trend estimate."
            )
    else:
        result["predictive"].append(
            "A time-based forecast is not available because a usable date/period and sales column were not both detected."
        )

    # Prescriptive recommendations.
    if result["sales_col"] and result["cost_col"]:
        result["prescriptive"].append(
            "Compare high-sales items with their unit costs and margins before increasing stock or marketing spend."
        )
    if result["missing"] > 0:
        result["prescriptive"].append(
            "Review missing cells in important business columns before making decisions."
        )
    if result["quality"] < 95:
        result["prescriptive"].append(
            "Clean inconsistent or incomplete values to improve analytical reliability."
        )
    if not result["prescriptive"]:
        result["prescriptive"].append(
            "Use the strongest products/categories as the first candidates for growth initiatives, while monitoring cost and margin."
        )

    return result


# ------------------------------------------------------------
# TEXT DOCUMENT ANALYSIS
# ------------------------------------------------------------
def analyze_text(text):
    text = text or ""
    clean = re.sub(r"\s+", " ", text).strip()
    words = re.findall(r"\b[\w₹$€£%.-]+\b", clean, flags=re.UNICODE)

    money_values = re.findall(
        r"(?:₹|\$|€|£)\s?\d[\d,]*(?:\.\d+)?",
        clean,
        flags=re.IGNORECASE,
    )
    percentages = re.findall(r"\b\d+(?:\.\d+)?\s?%", clean)

    # Sentence extraction.
    sentences = re.split(r"(?<=[.!?])\s+", clean)
    sentences = [s.strip() for s in sentences if len(s.strip()) > 25]

    # Business keywords to surface likely important sentences.
    business_terms = [
        "sales", "revenue", "profit", "loss", "cost", "customer", "market",
        "growth", "decline", "increase", "decrease", "risk", "target",
        "performance", "margin", "order", "quantity", "strategy", "recommendation",
    ]

    important = []
    for sentence in sentences:
        lower = sentence.lower()
        score = sum(1 for term in business_terms if term in lower)
        if score:
            important.append((score, sentence))
    important.sort(key=lambda x: x[0], reverse=True)

    key_points = []
    seen = set()
    for _, sentence in important[:8]:
        short = safe_text(sentence, 280)
        if short.lower() not in seen:
            key_points.append(short)
            seen.add(short.lower())

    if not key_points:
        key_points = [safe_text(s, 280) for s in sentences[:5]]

    return {
        "characters": len(clean),
        "words": len(words),
        "paragraphs": len([p for p in re.split(r"\n\s*\n", text) if p.strip()]),
        "money_values": money_values[:20],
        "percentages": percentages[:20],
        "key_points": key_points,
        "preview": clean,
    }


def find_text_answers(text, question):
    """Simple local Q&A without AI. It searches the uploaded text for relevant sentences."""
    clean = re.sub(r"\s+", " ", text or "").strip()
    if not clean:
        return ["There is no readable text in this document."]

    stopwords = {
        "what", "which", "when", "where", "who", "why", "how", "is", "are",
        "the", "a", "an", "of", "to", "for", "in", "on", "and", "or", "from",
        "with", "this", "that", "can", "could", "would", "should", "tell",
        "me", "about", "please", "does", "do", "did", "was", "were",
    }
    terms = [
        t.lower()
        for t in re.findall(r"[A-Za-z0-9₹$€£%.-]+", question)
        if t.lower() not in stopwords and len(t) > 2
    ]

    sentences = re.split(r"(?<=[.!?])\s+", clean)
    scored = []
    for sentence in sentences:
        lower = sentence.lower()
        score = sum(lower.count(term) for term in terms)
        if score:
            scored.append((score, sentence))

    scored.sort(key=lambda x: x[0], reverse=True)

    if scored:
        return [safe_text(s, 500) for _, s in scored[:5]]

    # Fallback for questions such as "summarize this document".
    if any(k in question.lower() for k in ["summary", "summarize", "overview", "key points"]):
        analysis = analyze_text(text)
        return analysis["key_points"][:5] or ["No clear key points were detected."]

    return [
        "I could not find a direct answer in the uploaded text. "
        "Try using a specific keyword from the document, such as a product, customer, month, sales figure, cost, or region."
    ]


def answer_table_question(df, question):
    q = question.lower().strip()
    if df.empty:
        return "The uploaded table is empty."

    analysis = analyze_table(df)
    metrics = analysis["metrics"]

    if any(k in q for k in ["total sales", "total revenue", "revenue total"]):
        if metrics["total_sales"] is not None:
            return f"Total sales/revenue is {money(metrics['total_sales'])}."

    if "total quantity" in q or "total units" in q:
        if metrics["total_quantity"] is not None:
            return f"Total quantity is {number(metrics['total_quantity'])}."

    if "average price" in q or "average unit price" in q:
        if metrics["average_price"] is not None:
            return f"Average price is {money(metrics['average_price'])}."

    if "profit" in q:
        if metrics["profit"] is not None:
            return f"Calculated total profit is {money(metrics['profit'])}."
        return "A profit value could not be calculated from the detected columns."

    if "margin" in q:
        if metrics["margin"] is not None:
            return f"Calculated margin is {percent(metrics['margin'])}."
        return "A margin could not be calculated from the detected sales and profit information."

    if any(k in q for k in ["missing", "blank", "empty cells"]):
        return f"The table contains {analysis['missing']:,} missing cells."

    if "rows" in q:
        return f"The table contains {analysis['rows']:,} rows."

    if "columns" in q:
        return f"The table contains {analysis['columns']:,} columns."

    # Ask about a particular product/category/column.
    for col in [analysis["product_col"], analysis["category_col"]]:
        if col:
            values = df[col].dropna().astype(str).unique().tolist()
            for value in values:
                if value.lower() in q:
                    subset = df[df[col].astype(str).str.lower() == value.lower()]
                    if analysis["sales_col"]:
                        sales = numeric_series(subset, analysis["sales_col"]).sum()
                        return f"{value} has {len(subset):,} records with total sales of {money(sales)}."
                    return f"{value} appears in {len(subset):,} records."

    # General search across text-like columns.
    terms = [t for t in re.findall(r"[A-Za-z0-9₹$€£.-]+", q) if len(t) > 2]
    text_cols = [c for c in df.columns if df[c].dtype == "object"][:8]
    matches = []
    for term in terms:
        for col in text_cols:
            mask = df[col].astype(str).str.contains(re.escape(term), case=False, na=False)
            if mask.any():
                matches.append((term, col, int(mask.sum())))
    if matches:
        term, col, count = matches[0]
        return f"I found {count:,} matching record(s) for '{term}' in the '{col}' column."

    return (
        "I can answer common questions about this table, including total sales, "
        "quantity, price, profit, margin, missing cells, products and categories. "
        "Try asking a more specific question."
    )


# ------------------------------------------------------------
# EXPORT
# ------------------------------------------------------------
def create_excel_report(df, analysis, source_name):
    """
    Uses XlsxWriter intentionally.
    This avoids the previous Streamlit Cloud crash caused by:
        pd.ExcelWriter(..., engine="openpyxl")
    """
    output = io.BytesIO()

    with pd.ExcelWriter(output, engine="xlsxwriter") as writer:
        workbook = writer.book

        title_fmt = workbook.add_format({
            "bold": True,
            "font_size": 16,
            "font_color": "white",
            "bg_color": "#1f4e78",
            "align": "left",
        })
        header_fmt = workbook.add_format({
            "bold": True,
            "font_color": "white",
            "bg_color": "#305496",
            "border": 1,
        })
        money_fmt = workbook.add_format({"num_format": '₹#,##0.00'})
        percent_fmt = workbook.add_format({"num_format": '0.0%'})
        wrap_fmt = workbook.add_format({"text_wrap": True, "valign": "top"})

        summary_rows = [
            ["Source file", source_name],
            ["Rows", analysis["rows"]],
            ["Columns", analysis["columns"]],
            ["Missing cells", analysis["missing"]],
            ["Quality score", analysis["quality"] / 100],
            ["Total Sales / Revenue", analysis["metrics"].get("total_sales")],
            ["Total Quantity", analysis["metrics"].get("total_quantity")],
            ["Average Price", analysis["metrics"].get("average_price")],
            ["Calculated Profit", analysis["metrics"].get("profit")],
            ["Calculated Margin", (
                analysis["metrics"].get("margin") / 100
                if analysis["metrics"].get("margin") is not None else None
            )],
        ]

        summary = pd.DataFrame(summary_rows, columns=["Metric", "Value"])
        summary.to_excel(writer, sheet_name="Summary", index=False, startrow=1)
        ws = writer.sheets["Summary"]
        ws.write(0, 0, "Universal Document Analytics Report", title_fmt)
        ws.set_column("A:A", 28)
        ws.set_column("B:B", 24)
        ws.set_column("A:B", 28)

        # Apply formats to values.
        for row_idx, (metric, value) in enumerate(summary_rows, start=2):
            if metric in {"Total Sales / Revenue", "Average Price", "Calculated Profit"}:
                ws.write(row_idx - 1, 1, value, money_fmt)
            elif metric == "Calculated Margin" or metric == "Quality score":
                ws.write(row_idx - 1, 1, value, percent_fmt)

        df.to_excel(writer, sheet_name="Data", index=False)
        data_ws = writer.sheets["Data"]
        data_ws.freeze_panes(1, 0)
        data_ws.autofilter(0, 0, max(0, len(df)), max(0, len(df.columns) - 1))
        for idx, col in enumerate(df.columns):
            width = min(40, max(12, len(str(col)) + 2))
            data_ws.set_column(idx, idx, width)

        highlight_rows = []
        for icon, title, label, value in analysis["highlights"]:
            highlight_rows.append([title, label, value])

        pd.DataFrame(
            highlight_rows,
            columns=["Highlight", "Item", "Value"]
        ).to_excel(writer, sheet_name="Highlights", index=False)
        hws = writer.sheets["Highlights"]
        hws.set_column("A:A", 24)
        hws.set_column("B:B", 28)
        hws.set_column("C:C", 28)

        for sheet_name, values in [
            ("Diagnostic", analysis["diagnostic"]),
            ("Predictive", analysis["predictive"]),
            ("Prescriptive", analysis["prescriptive"]),
        ]:
            pd.DataFrame({"Finding": values}).to_excel(
                writer, sheet_name=sheet_name, index=False
            )
            ws2 = writer.sheets[sheet_name]
            ws2.set_column("A:A", 100, wrap_fmt)

    return output.getvalue()


# ------------------------------------------------------------
# SIDEBAR / FILE UPLOAD
# ------------------------------------------------------------
st.sidebar.markdown("## 📂 Upload & Manage Files")
st.sidebar.caption("Upload a business document to analyze it.")

uploaded_file = st.sidebar.file_uploader(
    "Upload a document",
    type=[
        "xlsx", "xls", "xlsm", "csv", "json",
        "docx", "pdf", "pptx", "txt", "md",
        "html", "htm", "xml", "rtf",
    ],
    help="Supported: Excel, CSV, JSON, Word, PDF, PowerPoint, text and common markup files.",
)

if st.sidebar.button("🗑️ Clear current file", use_container_width=True):
    for key in ["file_result", "analysis", "question_answer"]:
        st.session_state.pop(key, None)
    st.rerun()

# ------------------------------------------------------------
# HEADER
# ------------------------------------------------------------
st.markdown('<div class="main-title">📊 Universal Document Analytics</div>', unsafe_allow_html=True)
st.markdown(
    '<div class="subtitle">Upload a business document and turn raw information into clear, useful insights.</div>',
    unsafe_allow_html=True,
)

if not uploaded_file:
    st.info(
        "👈 Upload a document from the sidebar to begin. "
        "The app works without external AI services."
    )

    st.markdown("### Supported file types")
    c1, c2, c3 = st.columns(3)
    with c1:
        st.markdown("**📊 Structured data**")
        st.write("Excel `.xlsx`, `.xls`, `.xlsm`")
        st.write("CSV `.csv`")
        st.write("JSON `.json`")
    with c2:
        st.markdown("**📄 Documents**")
        st.write("Word `.docx`")
        st.write("PDF `.pdf`")
        st.write("PowerPoint `.pptx`")
    with c3:
        st.markdown("**📝 Text**")
        st.write("TXT, Markdown, HTML, XML, RTF")

    st.stop()

# ------------------------------------------------------------
# READ FILE
# ------------------------------------------------------------
file_result = read_uploaded_file(uploaded_file)

if file_result["error"]:
    st.error(
        "We could not read this file. Please check that it is not damaged, "
        "password protected, or an unsupported format."
    )
    st.caption(f"Technical detail: {file_result['error']}")
    st.stop()

st.sidebar.success("File loaded successfully.")
st.sidebar.markdown(f"**Current file:** {uploaded_file.name}")
st.sidebar.markdown(f"**Type:** {'Table' if file_result['kind'] == 'table' else 'Text document'}")

# ------------------------------------------------------------
# TABLE DASHBOARD
# ------------------------------------------------------------
if file_result["kind"] == "table":
    df = file_result["dataframe"]
    analysis = analyze_table(df)

    st.session_state["analysis"] = analysis

    tabs = st.tabs([
        "📊 Overview",
        "🔎 Explore Data",
        "🧠 Analytics",
        "💬 Ask About Document",
        "📥 Export",
    ])

    with tabs[0]:
        st.subheader("Data Quality")

        q1, q2, q3, q4 = st.columns(4)
        cards = [
            ("Rows", number(analysis["rows"])),
            ("Columns", number(analysis["columns"])),
            ("Missing Cells", number(analysis["missing"])),
            ("Quality Score", percent(analysis["quality"])),
        ]
        for col, (label, value) in zip([q1, q2, q3, q4], cards):
            with col:
                st.markdown(
                    f'<div class="metric-card"><div class="metric-label">{label}</div>'
                    f'<div class="metric-value">{value}</div></div>',
                    unsafe_allow_html=True,
                )

        st.markdown("### Key Business Metrics")
        m1, m2, m3, m4 = st.columns(4)
        metric_cards = [
            ("Total Sales / Revenue", money(analysis["metrics"].get("total_sales"))),
            ("Total Quantity", number(analysis["metrics"].get("total_quantity"))),
            ("Average Price", money(analysis["metrics"].get("average_price"))),
            ("Calculated Profit", money(analysis["metrics"].get("profit"))),
        ]
        for col, (label, value) in zip([m1, m2, m3, m4], metric_cards):
            with col:
                st.markdown(
                    f'<div class="metric-card"><div class="metric-label">{label}</div>'
                    f'<div class="metric-value">{value}</div></div>',
                    unsafe_allow_html=True,
                )

        if analysis["highlights"]:
            st.markdown("### 🔥 Key Highlights")
            hcols = st.columns(min(3, len(analysis["highlights"])))
            for col, (icon, title, label, value) in zip(hcols, analysis["highlights"]):
                with col:
                    st.markdown(
                        f'<div class="insight-card">'
                        f'<div class="insight-title">{icon} {title}</div>'
                        f'<div style="font-size:20px;font-weight:700">{safe_text(label, 80)}</div>'
                        f'<div class="small-muted">{safe_text(value, 100)}</div>'
                        f'</div>',
                        unsafe_allow_html=True,
                    )

        detected = []
        mapping = [
            ("Sales/Revenue", analysis["sales_col"]),
            ("Quantity", analysis["quantity_col"]),
            ("Price", analysis["price_col"]),
            ("Cost", analysis["cost_col"]),
            ("Profit", analysis["profit_col"]),
            ("Product", analysis["product_col"]),
            ("Category", analysis["category_col"]),
            ("Period", analysis["period_col"]),
            ("Date", analysis["date_col"]),
        ]
        for label, col_name in mapping:
            if col_name:
                detected.append(f"**{label}:** `{col_name}`")

        if detected:
            st.markdown("### Detected Business Fields")
            st.write(" · ".join(detected))

    with tabs[1]:
        st.subheader("Data Preview")
        st.caption(f"Showing up to 100 rows from **{uploaded_file.name}**.")
        st.dataframe(df.head(100), use_container_width=True, height=430)

        st.subheader("Column Information")
        info = pd.DataFrame({
            "Column": df.columns,
            "Data Type": [str(df[c].dtype) for c in df.columns],
            "Non-Empty": [int(df[c].notna().sum()) for c in df.columns],
            "Missing": [int(df[c].isna().sum()) for c in df.columns],
        })
        st.dataframe(info, use_container_width=True, hide_index=True)

    with tabs[2]:
        st.subheader("📈 Analytics Type")

        sections = [
            (
                "Descriptive",
                "What happened?",
                [
                    f"The dataset contains {analysis['rows']:,} records across {analysis['columns']:,} columns.",
                    (
                        f"Total sales/revenue: {money(analysis['metrics']['total_sales'])}."
                        if analysis["metrics"]["total_sales"] is not None
                        else "No sales/revenue column was confidently detected."
                    ),
                    (
                        f"Total quantity: {number(analysis['metrics']['total_quantity'])}."
                        if analysis["metrics"]["total_quantity"] is not None
                        else "No quantity column was confidently detected."
                    ),
                ],
            ),
            (
                "Diagnostic",
                "Why?",
                analysis["diagnostic"]
                or ["Not enough recognizable product/category sales fields were found for a deeper comparison."],
            ),
            (
                "Predictive",
                "What next?",
                analysis["predictive"],
            ),
            (
                "Prescriptive",
                "What should we do?",
                analysis["prescriptive"],
            ),
        ]

        cols = st.columns(4)
        for col, (title, question, answers) in zip(cols, sections):
            with col:
                st.markdown(
                    f'<div class="insight-card"><div class="insight-title">{title}</div>'
                    f'<div class="small-muted">{question}</div></div>',
                    unsafe_allow_html=True,
                )
                for answer in answers:
                    st.write("• " + answer)

        st.markdown("### 📌 Detailed Findings")
        for label, values in [
            ("Descriptive", sections[0][2]),
            ("Diagnostic", analysis["diagnostic"]),
            ("Predictive", analysis["predictive"]),
            ("Prescriptive", analysis["prescriptive"]),
        ]:
            with st.expander(label):
                for value in values:
                    st.write("• " + value)

    with tabs[3]:
        st.subheader("💬 Ask About Your Document")
        st.caption(
            "This assistant works locally from the uploaded file. "
            "It does not send your document to an external AI service."
        )

        question = st.text_input(
            "Ask a question",
            placeholder="Example: What are the total sales? Which product performs best?",
            key="table_question",
        )

        if st.button("🔎 Answer Question", type="primary"):
            if not question.strip():
                st.warning("Please enter a question.")
            else:
                answer = answer_table_question(df, question)
                st.success(answer)

    with tabs[4]:
        st.subheader("📥 Export Analysis")

        report_bytes = create_excel_report(
            df,
            analysis,
            uploaded_file.name,
        )

        st.download_button(
            "⬇️ Download Complete Excel Report",
            data=report_bytes,
            file_name=f"{Path(uploaded_file.name).stem}_analytics_report.xlsx",
            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            use_container_width=True,
        )

        csv_bytes = df.to_csv(index=False).encode("utf-8-sig")
        st.download_button(
            "⬇️ Download Cleaned CSV",
            data=csv_bytes,
            file_name=f"{Path(uploaded_file.name).stem}_cleaned.csv",
            mime="text/csv",
            use_container_width=True,
        )

# ------------------------------------------------------------
# TEXT DOCUMENT DASHBOARD
# ------------------------------------------------------------
else:
    text = file_result["text"]
    text_analysis = analyze_text(text)

    tabs = st.tabs([
        "✨ Overview",
        "📝 Key Points",
        "📊 Document Statistics",
        "💬 Ask About Document",
    ])

    with tabs[0]:
        st.subheader("✨ Document Overview")

        c1, c2, c3, c4 = st.columns(4)
        cards = [
            ("Words", number(text_analysis["words"])),
            ("Characters", number(text_analysis["characters"])),
            ("Key Points", number(len(text_analysis["key_points"]))),
            ("Amounts Found", number(len(text_analysis["money_values"]))),
        ]
        for col, (label, value) in zip([c1, c2, c3, c4], cards):
            with col:
                st.markdown(
                    f'<div class="metric-card"><div class="metric-label">{label}</div>'
                    f'<div class="metric-value">{value}</div></div>',
                    unsafe_allow_html=True,
                )

        st.markdown("### 🔥 Key Business Highlights")
        if text_analysis["key_points"]:
            for i, point in enumerate(text_analysis["key_points"][:6], start=1):
                st.markdown(
                    f'<div class="insight-card" style="margin-bottom:10px">'
                    f'<div class="insight-title">🔹 Key Point {i}</div>'
                    f'<div>{safe_text(point, 500)}</div>'
                    f'</div>',
                    unsafe_allow_html=True,
                )
        else:
            st.info("No clear business key points were detected in the document.")

        if text_analysis["money_values"]:
            st.markdown("### 💰 Important Amounts Detected")
            st.write(" · ".join(text_analysis["money_values"][:15]))

        if text_analysis["percentages"]:
            st.markdown("### 📈 Percentages Detected")
            st.write(" · ".join(text_analysis["percentages"][:15]))

    with tabs[1]:
        st.subheader("📝 Key Points")
        for point in text_analysis["key_points"]:
            st.write("• " + point)

        st.markdown("### 📄 Document Preview")
        preview = text_analysis["preview"]
        st.write(safe_text(preview, 5000) if preview else "No readable text was extracted.")

    with tabs[2]:
        st.subheader("📊 Document Statistics")
        stats = pd.DataFrame({
            "Metric": ["Words", "Characters", "Paragraphs", "Amounts detected", "Percentages detected"],
            "Value": [
                text_analysis["words"],
                text_analysis["characters"],
                text_analysis["paragraphs"],
                len(text_analysis["money_values"]),
                len(text_analysis["percentages"]),
            ],
        })
        st.dataframe(stats, use_container_width=True, hide_index=True)

    with tabs[3]:
        st.subheader("💬 Ask About Document")
        st.caption(
            "Local document search is used here; no external AI service is required."
        )

        question = st.text_input(
            "Ask a question about this document",
            placeholder="Example: What does the document say about sales growth?",
            key="text_question",
        )

        if st.button("🔎 Answer Question", type="primary", key="text_answer_button"):
            if not question.strip():
                st.warning("Please enter a question.")
            else:
                answers = find_text_answers(text, question)
                st.markdown("### Answer")
                for answer in answers:
                    st.write("• " + answer)


st.caption(
    "Universal Document Analytics • Local analysis mode • "
    "Use only with documents and personal information you are authorized to process."
)
